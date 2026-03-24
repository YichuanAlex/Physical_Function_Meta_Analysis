import pandas as pd
import numpy as np
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

# Function to create three-line table style
def create_three_line_table(sheet, data_df):
    # Add data to sheet
    for r_idx, row in enumerate(dataframe_to_rows(data_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            sheet.cell(row=r_idx, column=c_idx, value=value)
    
    # Define borders
    thin_border = Side(style='thin', color='000000')
    thick_border = Side(style='thick', color='000000')
    
    # Apply top border to header row
    for cell in sheet[1]:
        cell.border = Border(top=thick_border, bottom=thin_border)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Apply bottom border to last data row
    last_row = sheet.max_row
    for cell in sheet[last_row]:
        cell.border = Border(bottom=thick_border)
    
    # Auto-adjust column widths
    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        sheet.column_dimensions[column_letter].width = adjusted_width

# Function to extract sub-tables from the main Excel file with correct boundaries
def extract_subtables():
    excel_path = '/Users/jiangzixi/Downloads/出图/CMA- analysis.xlsx'
    xl = pd.ExcelFile(excel_path)
    df = xl.parse('Sheet1')
    
    # Define table boundaries based on user's structure
    table_boundaries = {
        'grip_strength': {'start': 0, 'end': 10, 'outcome': 'grip strength'},
        '30CST': {'start': 12, 'end': 17, 'outcome': '30CST'},
        'FTSST': {'start': 19, 'end': 23, 'outcome': 'FTSST'},
        'muscle_mass': {'start': 24, 'end': 32, 'outcome': 'Muscle mass'},
        'MFES': {'start': 34, 'end': 37, 'outcome': 'MFES'},
        'gait_speed': {'start': 39, 'end': 42, 'outcome': 'Gait speed'},
        'TUG': {'start': 44, 'end': 49, 'outcome': 'TUG'}
    }
    
    # Create directory for sub-tables
    subtables_dir = '/Users/jiangzixi/Downloads/出图/subtables_academic'
    os.makedirs(subtables_dir, exist_ok=True)
    
    # Clean and save each sub-table
    for name, bounds in table_boundaries.items():
        print(f"\nProcessing {name}:")
        
        # Extract the specific rows for this table
        table = df.iloc[bounds['start']:bounds['end'], :9].copy()
        
        # Handle each table type specifically
        if name == 'gait_speed':
            # Special case for gait_speed
            table_clean = table.copy()
            
            # Set column names
            column_names = ['study', 'exp_mean', 'exp_sd', 'exp_n', 
                           'ctrl_mean', 'ctrl_sd', 'ctrl_n', 'sub']
            
            # Extract data rows (skip the first row which has numbers 1-6)
            data_rows = table_clean.iloc[1:].copy()
            data_rows.columns = column_names
            
            # Add outcome column
            data_rows['outcome'] = bounds['outcome']
            
        else:
            # Normal case
            table_clean = table.copy()
            
            # Extract data rows (skip header rows as needed)
            data_rows = table_clean.iloc[1:].copy()
            
            # Set column names
            column_names = ['study', 'exp_mean', 'exp_sd', 'exp_n', 
                           'ctrl_mean', 'ctrl_sd', 'ctrl_n', 'sub']
            data_rows.columns = column_names[:len(data_rows.columns)]
            
            # Add outcome column
            data_rows['outcome'] = bounds['outcome']
        
        # Reorder columns for consistency
        standard_columns = ['outcome', 'study', 'exp_mean', 'exp_sd', 'exp_n', 
                           'ctrl_mean', 'ctrl_sd', 'ctrl_n', 'sub']
        data_rows = data_rows.reindex(columns=standard_columns)
        
        # Drop any rows with missing critical data
        critical_columns = ['study', 'exp_mean', 'exp_sd', 'exp_n', 'ctrl_mean', 'ctrl_sd', 'ctrl_n']
        data_rows = data_rows.dropna(subset=critical_columns, how='any').reset_index(drop=True)
        
        # Save as three-line table in Excel
        output_path = os.path.join(subtables_dir, f'{name}.xlsx')
        
        # Create workbook and sheet
        wb = Workbook()
        ws = wb.active
        ws.title = name.replace('_', ' ').title()
        
        # Create three-line table
        create_three_line_table(ws, data_rows)
        
        # Save workbook
        wb.save(output_path)
        
        print(f"Saved {name} to {output_path}")
        print(f"Table shape: {data_rows.shape}")
        print(data_rows.head())
    
    return subtables_dir

# Main function
if __name__ == '__main__':
    extract_subtables()